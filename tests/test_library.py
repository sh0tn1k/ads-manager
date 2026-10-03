from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ads_api.models import Ad
from bot.formatters import ad_json
from bot.library import (
    LibraryStore,
    excel_sheet_name,
    flatten_value,
    normalize_title,
    snapshot_cells,
)


def _ad(**kwargs) -> Ad:
    payload = {
        "ad_id": 156,
        "title": "AI Posted",
        "currency": "TON",
        "promote_url": "https://t.me/ai_posted",
        "cpm": 0.1,
        "placement": "search_result",
        "spent_budget": 10,
        "remaining_budget": 1,
        "views": 1000,
        "clicks": 5,
        "actions": 2,
        "status": "active",
        "is_paused": False,
        "text": "Hello",
        "target": {"type": "search", "search_queries": ["vpn", "proxy"]},
    }
    payload.update(kwargs)
    return Ad.model_validate(payload)


def _store(name: str) -> LibraryStore:
    path = Path("data") / f"{name}.sqlite"
    path.parent.mkdir(exist_ok=True)
    for leftover in path.parent.glob(f"{name}.sqlite*"):
        leftover.unlink(missing_ok=True)
    return LibraryStore(path)


def test_flatten_and_sheet_names():
    out: dict = {}
    flatten_value({"target": {"type": "search", "search_queries": ["vpn", "proxy"]}}, "", out)
    assert out["target.type"] == "search"
    assert out["target.search_queries"] == "vpn, proxy"
    used: set[str] = set()
    assert excel_sheet_name("Группа 1", used) == "Группа 1"
    assert excel_sheet_name("a/b:c?*", used) == "a b c"
    assert excel_sheet_name("Группа 1", used) == "Группа 1_2"
    assert normalize_title("  поиск   апрель  ") == "поиск апрель"
    with pytest.raises(ValueError):
        normalize_title("   ")


@pytest.mark.asyncio
async def test_library_store_groups_and_upsert():
    store = _store("test_library_store")
    one = await store.create_group(user_id=1, chat_id=10, title="Группа 1")
    same = await store.create_group(user_id=1, chat_id=10, title="Группа 1")
    two = await store.create_group(user_id=1, chat_id=10, title="Группа 2")
    other = await store.create_group(user_id=2, chat_id=20, title="Группа 1")
    assert one.id == same.id
    assert two.id != one.id
    assert other.id != one.id

    await store.save_ad(group_id=one.id, user_id=1, ad_id=156, payload={"ad": ad_json(_ad())})
    await store.save_ad(
        group_id=one.id,
        user_id=1,
        ad_id=156,
        payload={"ad": ad_json(_ad(views=2000, title="Updated"))},
    )
    await store.save_ad(group_id=two.id, user_id=1, ad_id=200, payload={"ad": ad_json(_ad(ad_id=200))})

    counts = {group.title: group.ads_count for group in await store.list_groups(1)}
    assert counts["Группа 1"] == 1
    assert counts["Группа 2"] == 1
    other_groups = await store.list_groups(2)
    assert [group.title for group in other_groups] == ["Группа 1"]
    assert other_groups[0].id == other.id

    items = await store.list_ads(one.id, 1)
    assert len(items) == 1
    assert items[0].ad_title == "Updated"
    assert items[0].payload["ad"]["views"] == 2000
    assert await store.list_ads(one.id, 2) == []

    cells = snapshot_cells(items[0])
    assert cells["ad_id"] == 156
    assert cells["computed.cpc"] == 2.0
    assert cells["target.search_queries"] == "vpn, proxy"
    assert "Hello" in str(cells["json"])

    ok = await store.delete_ad(one.id, 1, 156)
    assert ok
    assert await store.list_ads(one.id, 1) == []
    await store.delete_group(two.id, 1)
    titles = [group.title for group in await store.list_groups(1)]
    assert "Группа 2" not in titles


@pytest.mark.asyncio
async def test_excel_workbook_one_sheet_per_group():
    store = _store("test_library_excel")
    one = await store.create_group(user_id=7, chat_id=1, title="Поиск апрель")
    two = await store.create_group(user_id=7, chat_id=1, title="Каналы")
    await store.save_ad(group_id=one.id, user_id=7, ad_id=156, payload={"ad": ad_json(_ad())})
    await store.save_ad(
        group_id=two.id,
        user_id=7,
        ad_id=9,
        payload={"ad": ad_json(_ad(ad_id=9, title="Other"))},
    )
    await store.create_group(user_id=7, chat_id=1, title="Пустая")

    data = await store.export_user(7)
    book = load_workbook(BytesIO(data))
    assert "Поиск апрель" in book.sheetnames
    assert "Каналы" in book.sheetnames
    assert "Пустая" in book.sheetnames
    sheet = book["Поиск апрель"]
    headers = [cell.value for cell in sheet[1]]
    assert "ID" in headers
    assert "Название" in headers
    assert "CTR" in headers
    assert "CPC" in headers
    assert headers[-1] == "JSON"
    id_col = headers.index("ID") + 1
    title_col = headers.index("Название") + 1
    json_col = headers.index("JSON") + 1
    assert sheet.cell(2, id_col).value == 156
    assert sheet.cell(2, title_col).value == "AI Posted"
    assert "promote_url" in sheet.cell(2, json_col).value
    assert book["Пустая"].max_row == 1
