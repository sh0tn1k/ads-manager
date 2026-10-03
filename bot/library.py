"""Local saved-ad groups. Not part of Telegram Ads API."""

from __future__ import annotations

import asyncio
import json
import re
import sqlite3
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

MAX_TITLE = 64
MAX_JSON_CELL = 32767

PREFERRED_COLUMNS = [
    "saved_at",
    "ad_id",
    "title",
    "status",
    "is_paused",
    "placement",
    "currency",
    "text",
    "promote_url",
    "button",
    "website_name",
    "cpm",
    "views",
    "clicks",
    "opens",
    "actions",
    "joins",
    "computed.ctr",
    "computed.cvr",
    "computed.cpc",
    "computed.cpa",
    "computed.ecpm",
    "computed.budget",
    "spent_budget",
    "remaining_budget",
    "daily_spent_budget",
    "daily_budget_limit",
    "impression_frequency",
    "show_userpic",
    "action_type",
    "conversion_event_id",
    "additional_info",
    "created_date",
    "activate_date",
    "deactivate_date",
    "photo.photo_id",
    "photo.photo_url",
    "video.video_id",
    "video.video_url",
    "website_photo_id",
    "website_photo.photo_id",
    "json",
]

COLUMN_TITLES = {
    "saved_at": "Сохранено",
    "ad_id": "ID",
    "title": "Название",
    "status": "Статус",
    "is_paused": "Пауза",
    "placement": "Таргет",
    "currency": "Валюта",
    "text": "Текст",
    "promote_url": "URL",
    "button": "Кнопка",
    "website_name": "Сайт",
    "cpm": "CPM ставка",
    "views": "Показы",
    "clicks": "Клики",
    "opens": "Открытия",
    "actions": "Действия",
    "joins": "Вступления",
    "computed.ctr": "CTR",
    "computed.cvr": "CVR",
    "computed.cpc": "CPC",
    "computed.cpa": "CPA",
    "computed.ecpm": "eCPM",
    "computed.budget": "Бюджет",
    "spent_budget": "Потрачено",
    "remaining_budget": "Остаток",
    "daily_spent_budget": "Дневной расход",
    "daily_budget_limit": "Дневной лимит",
    "impression_frequency": "Частота",
    "show_userpic": "Аватар",
    "action_type": "Конверсия",
    "conversion_event_id": "Pixel event",
    "additional_info": "Доп. инфо",
    "created_date": "Создано",
    "activate_date": "Активация",
    "deactivate_date": "Деактивация",
    "photo.photo_id": "Фото ID",
    "photo.photo_url": "Фото URL",
    "video.video_id": "Видео ID",
    "video.video_url": "Видео URL",
    "website_photo_id": "Фото сайта ID",
    "website_photo.photo_id": "Фото сайта",
    "json": "JSON",
}

_SHEET_BAD = re.compile(r"[:\\/?*\[\]]+")
_DATE_KEYS = {"saved_at", "created_date", "activate_date", "deactivate_date", "created_at"}


@dataclass(slots=True)
class SavedGroup:
    id: int
    user_id: int
    chat_id: int
    title: str
    created_at: int
    ads_count: int = 0


@dataclass(slots=True)
class SavedAd:
    id: int
    group_id: int
    ad_id: int
    ad_title: str
    saved_at: int
    payload: dict[str, Any]


def normalize_title(title: str) -> str:
    cleaned = " ".join(title.split())
    if not cleaned:
        raise ValueError("empty title")
    if len(cleaned) > MAX_TITLE:
        cleaned = cleaned[:MAX_TITLE].rstrip()
    return cleaned


def flatten_value(value: Any, prefix: str, out: dict[str, Any]) -> None:
    if isinstance(value, dict):
        if not value:
            if prefix:
                out[prefix] = None
            return
        for key, item in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            flatten_value(item, path, out)
        return
    if isinstance(value, list):
        if not value:
            out[prefix] = None
            return
        if all(not isinstance(item, (dict, list)) for item in value):
            out[prefix] = ", ".join("" if item is None else str(item) for item in value)
            return
        out[prefix] = json.dumps(value, ensure_ascii=False, default=str)
        return
    out[prefix] = _cell_value(prefix, value)


def snapshot_cells(item: SavedAd) -> dict[str, Any]:
    ad = item.payload.get("ad") if isinstance(item.payload.get("ad"), dict) else item.payload
    if not isinstance(ad, dict):
        ad = {}
    cells: dict[str, Any] = {}
    flatten_value(ad, "", cells)
    cells["saved_at"] = _cell_value("saved_at", item.saved_at)
    raw = json.dumps(ad, ensure_ascii=False, default=str)
    cells["json"] = raw[:MAX_JSON_CELL]
    return cells


def ordered_columns(rows: list[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    columns: list[str] = []
    for key in PREFERRED_COLUMNS:
        if any(key in row for row in rows):
            columns.append(key)
            seen.add(key)
    extras = sorted({key for row in rows for key in row if key not in seen})
    columns.extend(extras)
    if "json" in columns:
        columns = [key for key in columns if key != "json"] + ["json"]
    return columns


def column_title(key: str) -> str:
    return COLUMN_TITLES.get(key, key)


def excel_sheet_name(title: str, used: set[str]) -> str:
    cleaned = _SHEET_BAD.sub(" ", title)
    cleaned = " ".join(cleaned.split()) or "Группа"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate.casefold() in {name.casefold() for name in used}:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def build_workbook(groups: list[tuple[SavedGroup, list[SavedAd]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(wrap_text=True, vertical="center")
    used: set[str] = set()
    first = True
    for group, ads in groups:
        rows = [snapshot_cells(item) for item in ads]
        columns = ordered_columns(rows) or list(PREFERRED_COLUMNS)
        sheet = book.active if first else book.create_sheet()
        first = False
        sheet.title = excel_sheet_name(group.title, used)
        for col, key in enumerate(columns, start=1):
            cell = sheet.cell(1, col, column_title(key))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row_idx, row in enumerate(rows, start=2):
            for col, key in enumerate(columns, start=1):
                sheet.cell(row_idx, col, row.get(key))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 22
        for col, key in enumerate(columns, start=1):
            samples = [column_title(key)] + [str(row.get(key) or "")[:40] for row in rows[:30]]
            width = min(42, max(12, max(len(sample) for sample in samples) + 2))
            sheet.column_dimensions[get_column_letter(col)].width = width
    if first:
        sheet = book.active
        sheet.title = "Группы"
        sheet["A1"] = "Нет сохранённых групп"
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def _cell_value(key: str, value: Any) -> Any:
    if value is None:
        return None
    name = key.rsplit(".", 1)[-1]
    if name in _DATE_KEYS or name.endswith("_date"):
        if isinstance(value, (int, float)) and value > 1_000_000_000:
            return datetime.fromtimestamp(int(value), tz=timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _group(row: sqlite3.Row) -> SavedGroup:
    keys = row.keys()
    return SavedGroup(
        id=row["id"],
        user_id=row["user_id"],
        chat_id=row["chat_id"],
        title=row["title"],
        created_at=row["created_at"],
        ads_count=row["ads_count"] if "ads_count" in keys else 0,
    )


def _saved_ad(row: sqlite3.Row) -> SavedAd:
    payload = json.loads(row["payload"])
    if not isinstance(payload, dict):
        payload = {}
    return SavedAd(
        id=row["id"],
        group_id=row["group_id"],
        ad_id=row["ad_id"],
        ad_title=row["ad_title"] or "",
        saved_at=row["saved_at"],
        payload=payload,
    )


class LibraryStore:
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = asyncio.Lock()
        self._init()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA foreign_keys=ON")
        return conn

    def _init(self) -> None:
        with self._connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS library_groups (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    chat_id INTEGER NOT NULL,
                    title TEXT NOT NULL,
                    created_at INTEGER NOT NULL,
                    UNIQUE(user_id, title)
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS library_ads (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    group_id INTEGER NOT NULL,
                    ad_id INTEGER NOT NULL,
                    ad_title TEXT,
                    saved_at INTEGER NOT NULL,
                    payload TEXT NOT NULL,
                    UNIQUE(group_id, ad_id),
                    FOREIGN KEY(group_id) REFERENCES library_groups(id) ON DELETE CASCADE
                )
                """
            )
            conn.commit()

    async def list_groups(self, user_id: int) -> list[SavedGroup]:
        async with self._lock:
            return await asyncio.to_thread(self._list_groups, user_id)

    def _list_groups(self, user_id: int) -> list[SavedGroup]:
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT g.*, COUNT(a.id) AS ads_count
                FROM library_groups g
                LEFT JOIN library_ads a ON a.group_id = g.id
                WHERE g.user_id = ?
                GROUP BY g.id
                ORDER BY g.id DESC
                """,
                (user_id,),
            ).fetchall()
            return [_group(row) for row in rows]

    async def get_group(self, group_id: int, user_id: int) -> SavedGroup | None:
        async with self._lock:
            return await asyncio.to_thread(self._get_group, group_id, user_id)

    def _get_group(self, group_id: int, user_id: int) -> SavedGroup | None:
        with self._connect() as conn:
            row = conn.execute(
                """
                SELECT g.*, COUNT(a.id) AS ads_count
                FROM library_groups g
                LEFT JOIN library_ads a ON a.group_id = g.id
                WHERE g.id = ? AND g.user_id = ?
                GROUP BY g.id
                """,
                (group_id, user_id),
            ).fetchone()
            return _group(row) if row else None

    async def create_group(self, *, user_id: int, chat_id: int, title: str) -> SavedGroup:
        title = normalize_title(title)
        async with self._lock:
            group_id = await asyncio.to_thread(self._create_group, user_id, chat_id, title)
        group = await self.get_group(group_id, user_id)
        assert group is not None
        return group

    def _create_group(self, user_id: int, chat_id: int, title: str) -> int:
        now = int(time.time())
        with self._connect() as conn:
            try:
                cur = conn.execute(
                    """
                    INSERT INTO library_groups (user_id, chat_id, title, created_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (user_id, chat_id, title, now),
                )
                conn.commit()
                return int(cur.lastrowid)
            except sqlite3.IntegrityError:
                row = conn.execute(
                    "SELECT id FROM library_groups WHERE user_id=? AND title=?",
                    (user_id, title),
                ).fetchone()
                if row is None:
                    raise
                return int(row["id"])

    async def delete_group(self, group_id: int, user_id: int) -> bool:
        async with self._lock:
            return await asyncio.to_thread(self._delete_group, group_id, user_id)

    def _delete_group(self, group_id: int, user_id: int) -> bool:
        with self._connect() as conn:
            cur = conn.execute(
                "DELETE FROM library_groups WHERE id=? AND user_id=?",
                (group_id, user_id),
            )
            conn.commit()
            return cur.rowcount > 0

    async def save_ad(
        self,
        *,
        group_id: int,
        user_id: int,
        ad_id: int,
        payload: dict[str, Any],
    ) -> SavedAd:
        async with self._lock:
            saved_id = await asyncio.to_thread(self._save_ad, group_id, user_id, ad_id, payload)
        item = await asyncio.to_thread(self._get_ad, saved_id)
        assert item is not None
        return item

    def _save_ad(self, group_id: int, user_id: int, ad_id: int, payload: dict[str, Any]) -> int:
        group = self._get_group(group_id, user_id)
        if group is None:
            raise PermissionError("group not found")
        ad = payload.get("ad") if isinstance(payload.get("ad"), dict) else payload
        title = ""
        if isinstance(ad, dict):
            title = str(ad.get("title") or "")
        raw = json.dumps(payload, ensure_ascii=False, default=str)
        now = int(time.time())
        with self._connect() as conn:
            conn.execute(
                """
                INSERT INTO library_ads (group_id, ad_id, ad_title, saved_at, payload)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(group_id, ad_id) DO UPDATE SET
                    ad_title=excluded.ad_title,
                    saved_at=excluded.saved_at,
                    payload=excluded.payload
                """,
                (group_id, ad_id, title, now, raw),
            )
            conn.commit()
            row = conn.execute(
                "SELECT id FROM library_ads WHERE group_id=? AND ad_id=?",
                (group_id, ad_id),
            ).fetchone()
            return int(row["id"])

    def _get_ad(self, saved_id: int) -> SavedAd | None:
        with self._connect() as conn:
            row = conn.execute("SELECT * FROM library_ads WHERE id=?", (saved_id,)).fetchone()
            return _saved_ad(row) if row else None

    async def list_ads(self, group_id: int, user_id: int) -> list[SavedAd]:
        async with self._lock:
            return await asyncio.to_thread(self._list_ads, group_id, user_id)

    def _list_ads(self, group_id: int, user_id: int) -> list[SavedAd]:
        group = self._get_group(group_id, user_id)
        if group is None:
            return []
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT * FROM library_ads WHERE group_id=? ORDER BY saved_at DESC, id DESC",
                (group_id,),
            ).fetchall()
            return [_saved_ad(row) for row in rows]

    async def delete_ad(self, group_id: int, user_id: int, ad_id: int) -> bool:
        async with self._lock:
            return await asyncio.to_thread(self._delete_ad, group_id, user_id, ad_id)

    def _delete_ad(self, group_id: int, user_id: int, ad_id: int) -> bool:
        group = self._get_group(group_id, user_id)
        if group is None:
            return False
        with self._connect() as conn:
            cur = conn.execute(
                "DELETE FROM library_ads WHERE group_id=? AND ad_id=?",
                (group_id, ad_id),
            )
            conn.commit()
            return cur.rowcount > 0

    async def ad_group_ids(self, user_id: int, ad_id: int) -> set[int]:
        async with self._lock:
            return await asyncio.to_thread(self._ad_group_ids, user_id, ad_id)

    def _ad_group_ids(self, user_id: int, ad_id: int) -> set[int]:
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT a.group_id
                FROM library_ads a
                JOIN library_groups g ON g.id = a.group_id
                WHERE g.user_id=? AND a.ad_id=?
                """,
                (user_id, ad_id),
            ).fetchall()
            return {int(row["group_id"]) for row in rows}

    async def export_user(self, user_id: int) -> bytes:
        groups = await self.list_groups(user_id)
        packed: list[tuple[SavedGroup, list[SavedAd]]] = []
        for group in groups:
            packed.append((group, await self.list_ads(group.id, user_id)))
        return build_workbook(packed)
